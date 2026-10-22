"""Reads a flat-pattern DXF and works out what the laser has to cut, then prices it.

Geometry: every cuttable entity (lines, arcs, circles, polylines, splines, ellipses) is
turned into a path, flattened finely enough that lengths are accurate to well under 0.1%,
and loose segments are joined end to end into contours. Each contour is one pierce.
The biggest contour that is not inside another one is the outside of a part; anything
inside it is a hole or cutout.

Pricing: cut time = cut length / feed rate + pierces x pierce time, then machine time,
material, setup and handling are added from the rates table, which the shop edits.
"""
from __future__ import annotations

import io
import math
from dataclasses import dataclass, field

import ezdxf

CUT_TYPES = {"LINE", "ARC", "CIRCLE", "LWPOLYLINE", "POLYLINE", "SPLINE", "ELLIPSE"}
# Layers that customers use for things the laser should not cut
SKIP_LAYER_WORDS = ("BEND", "ETCH", "MARK", "SCRIBE", "DIM", "TEXT", "NOTE", "TITLE", "BORDER", "CENTER", "CENTRE", "HIDDEN")
INSUNITS_TO_INCH = {1: 1.0, 2: 12.0, 4: 1 / 25.4, 5: 1 / 2.54, 6: 1000 / 25.4}
UNIT_NAMES = {1: "inches", 2: "feet", 4: "millimetres", 5: "centimetres", 6: "metres"}


@dataclass
class Contour:
    points: list[tuple[float, float]]   # in inches
    closed: bool
    length: float = 0.0
    area: float = 0.0                    # unsigned, closed contours only
    role: str = ""                       # outside / hole / open


@dataclass
class PartGeometry:
    name: str
    units_note: str
    contours: list[Contour] = field(default_factory=list)
    cut_length: float = 0.0              # inches
    pierces: int = 0
    holes: int = 0
    parts_in_file: int = 0
    width: float = 0.0
    height: float = 0.0
    net_area: float = 0.0                # sq in, outsides minus holes
    smallest_hole: float | None = None   # diameter of the smallest round hole, inches
    entities_used: int = 0
    skipped_layers: list[str] = field(default_factory=list)
    duplicates_removed: int = 0
    dashed_skipped: int = 0
    gaps_closed: int = 0
    warnings: list[str] = field(default_factory=list)
    reference_lines: list[list[tuple[float, float]]] = field(default_factory=list)  # bend/etch lines, drawn but not cut
    source: str = "DXF"
    blocked: str = ""                    # set when the file can't be priced as it is, with the reason
    model_thickness: float | None = None # plate thickness read from a STEP model, inches
    ignored_items: int = 0               # PDF: dimensions, notes, title block and border left out
    drawing_width: float = 0.0           # PDF: part width as drawn before any scale correction


def _dist(a, b):
    return math.hypot(a[0] - b[0], a[1] - b[1])


def _poly_length(pts):
    return sum(_dist(pts[i], pts[i + 1]) for i in range(len(pts) - 1))


def _poly_area(pts):
    s = 0.0
    for i in range(len(pts) - 1):
        s += pts[i][0] * pts[i + 1][1] - pts[i + 1][0] * pts[i][1]
    return abs(s) / 2


def _inside(pt, poly):
    x, y = pt
    hit = False
    for i in range(len(poly) - 1):
        (x1, y1), (x2, y2) = poly[i], poly[i + 1]
        if (y1 > y) != (y2 > y) and x < (x2 - x1) * (y - y1) / (y2 - y1) + x1:
            hit = not hit
    return hit


def _pieces(e, scale):
    """Yields (points, exact length) for one entity, in inches. Lines, arcs and circles
    are measured exactly; polylines are split into their real lines and arcs; splines and
    ellipses are measured along very finely spaced points on the true curve."""
    kind = e.dxftype()
    if kind in ("LWPOLYLINE", "POLYLINE"):
        for sub in e.virtual_entities():
            yield from _pieces(sub, scale)
        return
    if kind == "LINE":
        a, b = e.dxf.start, e.dxf.end
        pts = [(a.x, a.y), (b.x, b.y)]
        length = math.dist(pts[0], pts[1])
    elif kind == "ARC":
        span = (e.dxf.end_angle - e.dxf.start_angle) % 360 or 360
        pts = [(v.x, v.y) for v in e.flattening(sagitta=0.0005)]
        length = e.dxf.radius * math.radians(span)
    elif kind == "CIRCLE":
        pts = [(v.x, v.y) for v in e.flattening(sagitta=0.0005)]
        length = 2 * math.pi * e.dxf.radius
    else:  # SPLINE, ELLIPSE
        pts = [(v.x, v.y) for v in e.flattening(distance=0.0001)]
        length = _poly_length(pts)
    if len(pts) >= 2 and length > 1e-9:
        yield [(x * scale, y * scale) for x, y in pts], length * scale


def _n(k, word):
    return f"{k} {word}" + ("" if k == 1 else "s")


def _skip_layer(layer: str) -> bool:
    up = layer.upper()
    return any(w in up for w in SKIP_LAYER_WORDS)


def read_dxf(data: bytes, name: str, units_override: str = "auto", join_tol_in: float = 0.002) -> PartGeometry:
    # recover.read copes with Windows line endings, old code pages, binary DXF and
    # the small structural faults that customer files often have
    from ezdxf import recover
    doc, _ = recover.read(io.BytesIO(data))
    msp = doc.modelspace()
    code = doc.header.get("$INSUNITS", 0)
    if units_override == "inches":
        scale, units_note = 1.0, "inches (set by you)"
    elif units_override == "millimetres":
        scale, units_note = 1 / 25.4, "millimetres (set by you)"
    elif code in INSUNITS_TO_INCH:
        scale, units_note = INSUNITS_TO_INCH[code], f"{UNIT_NAMES[code]} (from the file)"
    else:
        scale, units_note = 1.0, "inches (the file doesn't say, so inches were assumed)"

    geo = PartGeometry(name=name, units_note=units_note)
    pieces: list[list[tuple[float, float]]] = []
    skipped = set()
    for e in msp:
        if e.dxftype() not in CUT_TYPES:
            continue
        layer = e.dxf.get("layer", "0")
        if _skip_layer(layer):
            skipped.add(layer)
            geo.reference_lines.extend(pts for pts, _ in _pieces(e, scale))
            continue
        # SolidWorks, Inventor and others export bend lines on the cut layer but dashed
        # (HIDDEN for bend down, CENTER for bend up), so a dashed line is never a cut
        lt = e.dxf.get("linetype", "BYLAYER").upper()
        if lt == "BYLAYER":
            lt = doc.layers.get(layer).dxf.linetype.upper() if doc.layers.has_entry(layer) else "CONTINUOUS"
        if lt not in ("CONTINUOUS", "BYBLOCK"):
            geo.dashed_skipped += 1
            geo.reference_lines.extend(pts for pts, _ in _pieces(e, scale))
            continue
        geo.entities_used += 1
        pieces.extend(_pieces(e, scale))
    geo.skipped_layers = sorted(skipped)

    # Customer files often have the same line drawn twice; the laser would cut it twice
    seen, unique = set(), []
    for pts, length in pieces:
        mid = pts[len(pts) // 2]
        key = tuple(sorted([tuple(round(c, 4) for c in pts[0]), tuple(round(c, 4) for c in pts[-1])])) + (round(length, 4), round(mid[0], 3), round(mid[1], 3))
        if key in seen:
            geo.duplicates_removed += 1
            continue
        seen.add(key)
        unique.append((pts, length))

    geo.contours = _chain(unique, join_tol_in, geo)
    _classify(geo)
    return geo


def _chain(pieces, tol, geo):
    contours, open_pieces = [], []
    for pts, length in pieces:
        if _dist(pts[0], pts[-1]) <= tol and len(pts) > 2:
            contours.append(Contour(points=pts[:-1] + [pts[0]], closed=True, length=length))
        else:
            open_pieces.append((list(pts), length))
    # join loose segments end to end, nearest match first
    while open_pieces:
        chain, chain_len = open_pieces.pop(0)
        grown = True
        while grown:
            grown = False
            if _dist(chain[0], chain[-1]) <= tol and len(chain) > 2:
                break
            best = None
            for i, (p, plen) in enumerate(open_pieces):
                for end_at_tail in (True, False):
                    anchor = chain[-1] if end_at_tail else chain[0]
                    for rev in (False, True):
                        q = p[::-1] if rev else p
                        d = _dist(anchor, q[0] if end_at_tail else q[-1])
                        if d <= tol and (best is None or d < best[0]):
                            best = (d, i, end_at_tail, q, plen)
            if best:
                d, i, end_at_tail, q, plen = best
                open_pieces.pop(i)
                if d > 1e-6:
                    geo.gaps_closed += 1
                chain = chain + q[1:] if end_at_tail else q[:-1] + chain
                chain_len += plen
                grown = True
        closed = _dist(chain[0], chain[-1]) <= tol and len(chain) > 2
        if closed:
            if _dist(chain[0], chain[-1]) > 1e-6:
                geo.gaps_closed += 1
            chain = chain[:-1] + [chain[0]]
        contours.append(Contour(points=chain, closed=closed, length=chain_len))
    for c in contours:
        c.area = _poly_area(c.points) if c.closed else 0.0
    return contours


def _classify(geo: PartGeometry):
    closed = sorted([c for c in geo.contours if c.closed], key=lambda c: -c.area)
    for i, c in enumerate(closed):
        depth = sum(1 for o in closed[:i] if o.area > c.area and _inside(c.points[0], o.points))
        c.role = "outside" if depth % 2 == 0 else "hole"
    for c in geo.contours:
        if not c.closed:
            c.role = "open"
    geo.cut_length = sum(c.length for c in geo.contours)
    geo.pierces = len(geo.contours)
    geo.holes = sum(1 for c in geo.contours if c.role == "hole")
    geo.parts_in_file = sum(1 for c in geo.contours if c.role == "outside")
    geo.net_area = sum(c.area for c in geo.contours if c.role == "outside") - sum(c.area for c in geo.contours if c.role == "hole")
    allpts = [p for c in geo.contours for p in c.points]
    if allpts:
        xs, ys = [p[0] for p in allpts], [p[1] for p in allpts]
        geo.width, geo.height = max(xs) - min(xs), max(ys) - min(ys)
    # a hole counts as round when its area matches a circle of the same perimeter
    round_holes = [c for c in geo.contours if c.role == "hole" and c.length > 0 and abs(4 * math.pi * c.area / c.length ** 2 - 1) < 0.01]
    if round_holes:
        geo.smallest_hole = min(c.length / math.pi for c in round_holes)

    open_n = sum(1 for c in geo.contours if c.role == "open")
    if not geo.contours:
        geo.warnings.append("No cuttable geometry found. Check the file isn't a block or only text.")
    if open_n:
        geo.warnings.append(f"{_n(open_n, 'open contour')}: the outline has a gap or there is a stray line. Each one is still counted as a cut and a pierce.")
    if geo.duplicates_removed:
        geo.warnings.append(f"{_n(geo.duplicates_removed, 'duplicate line')} removed so the laser doesn't cut the same path twice.")
    if geo.gaps_closed:
        geo.warnings.append(f"{_n(geo.gaps_closed, 'tiny gap')} closed when joining the outline.")
    if geo.parts_in_file > 1:
        geo.warnings.append(f"This file holds {geo.parts_in_file} separate parts. Totals cover all of them.")
    if geo.dashed_skipped:
        geo.warnings.append(f"{_n(geo.dashed_skipped, 'dashed line')} left out as bend or centre lines.")
    if geo.skipped_layers:
        geo.warnings.append("Left out because the layer name says bend, etch or notes: " + ", ".join(geo.skipped_layers))


def estimate(geo: PartGeometry, rate: dict, qty: int, settings: dict) -> dict:
    """rate: one row of the rates table. settings: shop-wide charges."""
    feed = float(rate["feed_ipm"])
    cut_min = geo.cut_length / feed if feed > 0 else 0.0
    pierce_min = geo.pierces * float(rate["pierce_sec"]) / 60
    run_min_each = (cut_min + pierce_min) * (1 + settings["motion_allowance_pct"] / 100)
    machine_each = run_min_each / 60 * settings["machine_rate_hr"]

    t = float(rate["thickness_in"])
    lb_cost = float(rate["density_lb_in3"]) * float(rate["price_per_lb"]) * (1 + settings["material_markup_pct"] / 100)
    sw, sh = settings.get("sheet_w_in", 48.0), settings.get("sheet_h_in", 96.0)
    per_sheet = sheet_fit(geo.width, geo.height, settings["part_spacing_in"], sw, sh, settings.get("sheet_margin_in", 0.5))
    sheets = math.ceil(qty / per_sheet) if per_sheet else 0
    basis = settings["material_basis"]
    if basis == "Full sheets used" and per_sheet:
        material_each = sheets * sw * sh * t * lb_cost / qty
    elif basis == "Net part area":
        material_each = geo.net_area * t * lb_cost
    else:
        material_each = (geo.width + settings["part_spacing_in"]) * (geo.height + settings["part_spacing_in"]) * t * lb_cost
    weight_each = geo.net_area * t * float(rate["density_lb_in3"])

    each = machine_each + material_each + settings["handling_per_part"]
    subtotal = each * qty + settings["setup_per_job"]
    total = max(subtotal, settings["minimum_charge"])
    warnings = []
    if not per_sheet:
        warnings.append(f"The part doesn't fit on a {sw:g} x {sh:g} in sheet.")
    if geo.smallest_hole is not None and geo.smallest_hole < t:
        warnings.append(f"Smallest hole ({geo.smallest_hole:.3f} in) is smaller than the material thickness ({t:.3f} in). It may need drilling or a special cut.")
    return {
        "cut_min_each": cut_min, "pierce_min_each": pierce_min, "run_min_each": run_min_each,
        "machine_each": machine_each, "material_each": material_each, "weight_each": weight_each,
        "each": each, "qty": qty, "setup": settings["setup_per_job"], "subtotal": subtotal,
        "total": total, "minimum_applied": total > subtotal, "warnings": warnings,
        "per_sheet": per_sheet, "sheets": sheets, "sheet": (sw, sh),
    }


def sheet_fit(w, h, spacing, sheet_w, sheet_h, margin=0.5) -> int:
    """Parts per sheet laid out in a plain grid, trying both ways round. Real nesting
    software packs tighter, so this is the safe side for a quote."""
    uw, uh = sheet_w - 2 * margin + spacing, sheet_h - 2 * margin + spacing
    pw, ph = w + spacing, h + spacing
    if pw <= 0 or ph <= 0:
        return 0
    return max(int(uw // pw) * int(uh // ph), int(uw // ph) * int(uh // pw))


def job_total(lines: list[dict], settings: dict) -> dict:
    """Setup is charged once per job and the minimum charge applies to the whole job,
    not to each part."""
    parts = sum(l["each"] * l["qty"] for l in lines)
    subtotal = parts + settings["setup_per_job"]
    total = max(subtotal, settings["minimum_charge"]) if lines else 0.0
    return {"parts": parts, "setup": settings["setup_per_job"], "subtotal": subtotal,
            "total": total, "minimum_applied": bool(lines) and total > subtotal,
            "run_min": sum(l["run_min_each"] * l["qty"] for l in lines)}


def svg_preview(geo: PartGeometry, size: int = 360) -> str:
    """Drawn the way shop quoting tools show a part: the plate filled light grey with a
    thin outline, holes cut out, bend and etch lines dashed, open outlines in red."""
    if not geo.contours:
        return ""
    allpts = [p for c in geo.contours for p in c.points]
    minx, miny = min(p[0] for p in allpts), min(p[1] for p in allpts)
    w, h = max(geo.width, 1e-6), max(geo.height, 1e-6)
    pad = max(w, h) * 0.04
    vb = f"{minx - pad:.4f} {-(miny + h) - pad:.4f} {w + 2 * pad:.4f} {h + 2 * pad:.4f}"
    stroke = max(w, h) / 420

    def d(points, close=False):
        return " ".join(f"{'M' if i == 0 else 'L'}{x:.4f},{-y:.4f}" for i, (x, y) in enumerate(points)) + (" Z" if close else "")

    closed = " ".join(d(c.points, True) for c in geo.contours if c.closed)
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{vb}" width="100%" style="max-height:{size}px;display:block" role="img" aria-label="Flat pattern of {geo.name}">']
    if closed:
        out.append(f'<path d="{closed}" fill="#e4e5e4" fill-rule="evenodd" stroke="#3b4045" stroke-width="{stroke:.4f}" stroke-linejoin="round"/>')
    for ref in geo.reference_lines:
        out.append(f'<path d="{d(ref)}" fill="none" stroke="#0e9aa7" stroke-width="{stroke:.4f}" stroke-dasharray="{stroke * 6:.4f} {stroke * 4:.4f}"/>')
    for c in geo.contours:
        if not c.closed:
            out.append(f'<path d="{d(c.points)}" fill="none" stroke="#c81e1e" stroke-width="{stroke * 1.6:.4f}"/>')
    out.append("</svg>")
    return "".join(out)


def read_step_file(data: bytes, name: str) -> PartGeometry:
    import step_reader
    sp = step_reader.read_step(data, name)
    geo = PartGeometry(name=name, units_note=sp.units_note, source="STEP")
    if not sp.flat:
        geo.blocked = sp.reason
        return geo
    geo.model_thickness = sp.thickness
    geo.contours = [Contour(points=pts, closed=True, length=length, area=_poly_area(pts)) for pts, length in zip(sp.loops, sp.loop_lengths)]
    _classify(geo)
    geo.units_note = f"{sp.units_note}; a flat plate {sp.thickness:.4f} in thick, read from the 3D model"
    return geo


def read_pdf(data: bytes, name: str, width_in: float | None = None) -> PartGeometry:
    """A vector PDF exported from CAD. The page border is dropped, the part's outline is
    the biggest closed shape left, and only closed shapes inside it are kept, so
    dimensions, centre marks, notes and the title block don't get priced."""
    import pymupdf
    geo = PartGeometry(name=name, units_note="drawn scale (1:1 unless you enter the real width)", source="PDF")
    doc = pymupdf.open(stream=data, filetype="pdf")
    page = doc[0]
    drawings = page.get_drawings()
    if not drawings:
        geo.blocked = ("This PDF has no lines inside it, only a picture (probably a scan), so nothing can be measured. "
                       "Ask the customer for the DXF, or the PDF exported straight from their CAD.")
        return geo
    pw, ph = page.rect.width, page.rect.height
    pieces, refs = [], []

    def pt(p):
        return (p.x / 72.0, (ph - p.y) / 72.0)

    for d in drawings:
        if d.get("color") is None:          # filled only: arrowheads, hatching, text drawn as shapes
            continue
        dashed = d.get("dashes") not in (None, "", "[] 0", "[] 0.0")
        for it in d["items"]:
            kind = it[0]
            if kind == "l":
                segs = [[pt(it[1]), pt(it[2])]]
            elif kind == "c":
                p0, c1, c2, p3 = (pt(q) for q in it[1:5])
                segs = [[(
                    (1 - t) ** 3 * p0[0] + 3 * (1 - t) ** 2 * t * c1[0] + 3 * (1 - t) * t * t * c2[0] + t ** 3 * p3[0],
                    (1 - t) ** 3 * p0[1] + 3 * (1 - t) ** 2 * t * c1[1] + 3 * (1 - t) * t * t * c2[1] + t ** 3 * p3[1])
                    for t in (i / 48 for i in range(49))]]
            elif kind == "re":
                r = it[1]
                c = [(r.x0 / 72, (ph - r.y0) / 72), (r.x1 / 72, (ph - r.y0) / 72), (r.x1 / 72, (ph - r.y1) / 72), (r.x0 / 72, (ph - r.y1) / 72)]
                segs = [[c[i], c[(i + 1) % 4]] for i in range(4)]
            elif kind == "qu":
                q = it[1]
                c = [pt(q.ul), pt(q.ur), pt(q.lr), pt(q.ll)]
                segs = [[c[i], c[(i + 1) % 4]] for i in range(4)]
            else:
                continue
            for seg in segs:
                (refs if dashed else pieces).append((seg, _poly_length(seg)))

    contours = _chain(pieces, 0.007, geo)
    geo.gaps_closed = 0  # joins in a PDF are rounding, not gaps in the drawing
    def is_border(c):
        # a drawing border is a plain rectangle running close to the page edges
        xs, ys = [x for x, _ in c.points], [y for _, y in c.points]
        x0, x1, y0, y1 = min(xs), max(xs), min(ys), max(ys)
        on_box = all(min(abs(x - x0), abs(x - x1), abs(y - y0), abs(y - y1)) < 0.01 for x, y in c.points)
        near_edges = x0 < 1.0 and y0 < 1.0 and x1 > pw / 72 - 1.0 and y1 > ph / 72 - 1.0
        return c.closed and on_box and near_edges
    frame = [c for c in contours if is_border(c)]
    rest = [c for c in contours if c not in frame and c.closed]
    if not rest:
        geo.blocked = "No closed outline found on the page, so the part can't be picked out of the drawing."
        return geo
    outline = max(rest, key=lambda c: c.area)
    kept = [outline] + [c for c in rest if c is not outline and c.area < outline.area and _inside(c.points[0], outline.points)]
    geo.ignored_items = len(contours) - len(kept)
    ox = [x for x, _ in outline.points]
    geo.drawing_width = max(ox) - min(ox)
    k = (width_in / geo.drawing_width) if width_in and geo.drawing_width else 1.0
    if k != 1.0:
        geo.units_note = f"scaled so the part is {width_in:.3f} in wide, as you entered (the print is drawn at {1 / k:.3f} of full size)"
        for c in kept:
            c.points = [(x * k, y * k) for x, y in c.points]
            c.length *= k
            c.area *= k * k
    for seg, _ in refs:
        oxmin, oxmax = min(ox), max(ox)
        if all(oxmin - 0.01 <= x <= oxmax + 0.01 for x, _ in seg):
            geo.reference_lines.append([(x * k, y * k) for x, y in seg])
    geo.contours = kept
    _classify(geo)
    if geo.ignored_items:
        geo.warnings.append(f"{_n(geo.ignored_items, 'item')} on the page left out (border, title block, dimensions, notes). Check the drawing above shows only the part.")
    return geo


RATE_COLUMNS = {"material": "material", "thickness": "thickness", "thickness (in": "thickness_in",
                "cut speed": "feed_ipm", "pierce time": "pierce_sec", "density": "density_lb_in3", "material price": "price_per_lb"}
CHARGE_KEYS = {"machine rate": "machine_rate_hr", "head travel": "motion_allowance_pct", "setup per job": "setup_per_job",
               "minimum job": "minimum_charge", "handling per part": "handling_per_part", "material markup": "material_markup_pct",
               "spacing between": "part_spacing_in", "sheet width": "sheet_w_in", "sheet length": "sheet_h_in",
               "sheet edge": "sheet_margin_in", "quote valid": "valid_days", "lead time": "lead_days"}


def load_rates_workbook(data: bytes) -> tuple[list[dict], dict]:
    """Reads the shop's rates.xlsx (Rates and Charges sheets). Columns are matched by the
    start of their heading, so small wording changes in Excel don't break it."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["Rates"]
    heads = [str(c.value or "").strip().lower() for c in ws[1]]

    def key_for(h):
        if h in ("material", "thickness"):
            return h
        # longest prefix wins, so "material price" isn't read as "material"
        hits = [k for k in RATE_COLUMNS if h.startswith(k) and k not in ("material", "thickness")]
        return RATE_COLUMNS[max(hits, key=len)] if hits else None
    keys = [key_for(h) for h in heads]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = {k: v for k, v in zip(keys, r) if k}
        if row.get("material") and row.get("thickness") is not None:
            row["thickness"] = str(row["thickness"])
            rows.append(row)
    settings = {}
    if "Charges" in wb.sheetnames:
        for label, value, *_ in wb["Charges"].iter_rows(min_row=2, values_only=True):
            key = next((v for k, v in CHARGE_KEYS.items() if str(label or "").lower().startswith(k)), None)
            if key and value is not None:
                settings[key] = float(value)
    return rows, settings
